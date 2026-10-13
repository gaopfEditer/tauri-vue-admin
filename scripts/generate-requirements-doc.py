# -*- coding: utf-8 -*-
"""从内置需求数据生成 docs/软件需求描述.md 与 docs/软件需求描述.xlsx"""
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

# 大模块 -> [(功能点, [细则条目...]), ...]
MODULES: dict[str, list[tuple[str, list[str]]]] = {
    "登录与授权": [
        (
            "登录_显示信息",
            [
                "登录页展示系统 Logo（可填充主题色）与软件名称（环境变量 VITE_APP_TITLE，默认可配置）；",
                "页面提供主题色渐变背景与 LoginBg 装饰背景，支持明暗主题切换；",
                "登录卡片内展示当前登录子模块标题（密码登录、验证码登录等）；",
                "建议扩展：公司名称、软件型号/版本号可在登录页固定区域展示（与授权或厂区 branding 配置联动）。",
            ],
        ),
        (
            "登录_密码登录",
            [
                "输入用户名、密码，支持「记住我」与「忘记密码」跳转；",
                "支持快捷切换验证码登录、注册账号模块；",
                "支持从历史/其他账号列表一键填充并登录；",
                "登录过程展示 loading，成功后进入系统首页（VITE_ROUTE_HOME_PATH）。",
            ],
        ),
        (
            "登录_多模块",
            [
                "路由支持 /login/:module，含 pwd-login、code-login、register、reset-pwd、bind-wechat；",
                "各子模块独立表单与校验规则；",
                "未授权访问业务页时重定向登录。",
            ],
        ),
        (
            "授权_锁定页",
            [
                "License 无效时进入独立授权页，展示校验原因、本机 MAC、绑定 MAC、有效期、客户信息；",
                "支持粘贴厂商授权码（PN1...）并激活；",
                "支持刷新授权状态，激活成功后跳转登录页。",
            ],
        ),
        (
            "授权_管理",
            [
                "已登录管理员可查看当前授权状态、剩余天数、设备锁与 MAC 绑定情况；",
                "支持导入激活码并绑定本机；",
                "厂商侧可使用 MASTER KEY 签发新授权（客户、生效/到期、是否设备锁）。",
            ],
        ),
    ],
    "仪表盘": [
        (
            "分析页_概览指标",
            [
                "展示在线传感器数量、未确认报警数、运行中采样数、报表数量等核心 KPI；",
                "按传感器类别（粒子/生物/模拟量）统计分布；",
                "展示近 7 日报警趋势折线图。",
            ],
        ),
        (
            "分析页_明细列表",
            [
                "底部区域展示近期报警与采样任务摘要列表；",
                "数据来源于传感器、报警、采样、报表、审计等聚合接口。",
            ],
        ),
        (
            "工作台",
            [
                "提供快捷入口与项目信息展示（模板能力，可配置业务快捷方式）。",
            ],
        ),
    ],
    "配方管理": [
        (
            "配方_列表与维护",
            [
                "配方列表分页展示名称、描述、关联传感器组、创建人、更新时间；",
                "支持新建、编辑、删除配方；",
                "编辑项含：名称、描述、传感器组多选、Pens(JSON) 趋势笔配置。",
            ],
        ),
        (
            "配方_关联",
            [
                "配方可被采样调度、实时趋势会话引用；",
                "传感器组数据来自传感器分组接口。",
            ],
        ),
    ],
    "采样管理": [
        (
            "采样_调度编辑器",
            [
                "列表展示配方、计划时间、模式、状态、自定义字段、备注等；",
                "支持调度新采样：选择配方、计划时间、采样模式（Operational/At Rest/Normal Operation/Custom）；",
                "动态渲染已启用的自定义字段（下拉或文本）；",
                "支持编辑、删除、中止（abort）运行中或计划任务；",
                "状态标签区分 scheduled / running / finished / aborted / failed。",
            ],
        ),
        (
            "采样_自定义字段",
            [
                "配置采样表单的扩展字段：Key、显示名、启用、可编辑、默认选项列表；",
                "支持增删行并保存至服务端，供采样编辑器动态使用。",
            ],
        ),
        (
            "采样_实时信号",
            [
                "按厂区房间树选择监控范围，卡片展示各设备实时指标与状态（正常/报警/离线等）；",
                "支持单设备勾选、批量开启/关闭电源；",
                "展示设备编码、房间、通道、多指标数值与通信/流量报警标识；",
                "支持刷新全部信号数据。",
            ],
        ),
    ],
    "设备传感器": [
        (
            "传感器_分类面板",
            [
                "粒子、生物、模拟量三类设备共用面板：卡片展示编号、描述、通道、最新值、运行状态；",
                "区分「在用 / 已删除」列表；",
                "支持新增设备、电源 ON/OFF、删除与恢复；",
                "粒子设备额外支持 Cleaning 模式与设备配置弹窗。",
            ],
        ),
        (
            "传感器_粒子配置",
            [
                "粒子设备专用配置模态框，保存 Modbus/采样相关参数（保存后刷新列表并记审计）。",
            ],
        ),
        (
            "传感器_统一编辑",
            [
                "传感器编辑器页面对设备进行跨类别检索与批量属性维护。",
            ],
        ),
        (
            "传感器_限值",
            [
                "按类别筛选传感器，配置数据类型、预警限、报警限及变更原因；",
                "支持新增/编辑限值、查看限值变更历史（含电子签名前后快照）；",
                "限值超限触发报警中心事件。",
            ],
        ),
    ],
    "报警中心": [
        (
            "报警_列表",
            [
                "展示报警时间、传感器、级别（warning/alarm/communication/flow）、消息、数据类型、当前值、限值；",
                "统计未确认数量与合计；",
                "未确认行高亮，已确认显示成功标签。",
            ],
        ),
        (
            "报警_确认",
            [
                "单条确认记录操作人与原因；",
                "支持「全部确认」批量处理未确认报警。",
            ],
        ),
    ],
    "实时趋势": [
        (
            "趋势_会话",
            [
                "选择配方后启动实时趋势会话，展示运行/停止状态与会话 ID；",
                "停止会话后结束轮询；",
                "进入页面时恢复当前活动会话状态。",
            ],
        ),
        (
            "趋势_数据点",
            [
                "运行中定时拉取趋势点表：时间、传感器、数据类型、数值；",
                "配方 Pens 配置决定采集笔定义。",
            ],
        ),
    ],
    "运行逻辑": [
        (
            "逻辑_Tags",
            [
                "展示 Available Tags 列表，供规则条件引用。",
            ],
        ),
        (
            "逻辑_Rules",
            [
                "规则列表含名称、类型、启用状态；",
                "支持新建/编辑/删除规则：条件 JSON、动作 JSON；",
                "支持测试输出（test output）验证规则执行结果。",
            ],
        ),
    ],
    "数据报表": [
        (
            "报表_生成器",
            [
                "选择报告类型：Audit Trail / Data / Trend / Sampling；",
                "选择模式 By Group / By Sampling，填写起止时间与采样间隔；",
                "生成后进入历史列表，支持预览与导出 CSV（记审计）。",
            ],
        ),
        (
            "报表_采样报告",
            [
                "针对采样任务生成专用报告视图与导出能力。",
            ],
        ),
        (
            "报表_Tags目录",
            [
                "维护与浏览报表/运行逻辑使用的 Tags 目录。",
            ],
        ),
        (
            "报表_统计分析器",
            [
                "SDA 统计分析：对历史数据做统计图表与分析（与后端 SDA 服务对接）。",
            ],
        ),
    ],
    "系统管理": [
        (
            "管理_用户",
            [
                "用户 CRUD、角色绑定、列设置等表格管理能力。",
            ],
        ),
        (
            "管理_角色",
            [
                "角色定义与权限集合配置。",
            ],
        ),
        (
            "管理_权限",
            [
                "权限项与角色/用户授权关系维护。",
            ],
        ),
        (
            "管理_路由",
            [
                "动态路由/菜单配置（与 VITE_AUTH_ROUTE_MODE 配合）。",
            ],
        ),
        (
            "管理_密码策略",
            [
                "配置密码过期天数、最小长度、历史密码不可复用次数、自动登出秒数；",
                "开关电子签名；",
                "展示用户组功能矩阵（Table 11-2）各角色可用功能标签。",
            ],
        ),
    ],
    "审计轨迹": [
        (
            "审计_查询",
            [
                "按类别、动作关键词、对象类型、操作人、时间范围筛选；",
                "列表展示时间、类别、动作、对象、操作人、详情摘要。",
            ],
        ),
        (
            "审计_自动记录",
            [
                "自动记录登录/登出、打开页面、配置保存、导出、传感器/配方/采样写操作、报警/趋势、系统运维等；",
                "支持操作员手动提交 User Event（事件描述与原因）。",
            ],
        ),
    ],
    "系统运维": [
        (
            "运维_备份恢复",
            [
                "配置日备开关、执行时间、备份路径、保留天数；",
                "立即备份与备份任务列表（类型、状态、文件路径、操作人、消息）。",
            ],
        ),
        (
            "运维_系统控制",
            [
                "高风险命令：关闭应用、重启、关机、断开客户端关机；",
                "执行前必须通过电子签名模态框确认，并写入审计。",
            ],
        ),
        (
            "运维_语言",
            [
                "系统界面语言偏好设置（与 i18n/后端配置联动）。",
            ],
        ),
        (
            "运维_缓冲重置",
            [
                "重置数据采集/通信缓冲，用于现场联调恢复。",
            ],
        ),
        (
            "运维_厂区组织",
            [
                "三级树：厂房/园区 → 区域/洁净区 → 房间；",
                "节点含编码、名称、描述、排序、启用状态；",
                "支持增删改与子节点创建，供实时信号等按房间筛选。",
            ],
        ),
        (
            "运维_设备轮询",
            [
                "Modbus 轮询联调：查看轮询日志与设备通信状态（MODBUS_POLL_* 环境配置）。",
            ],
        ),
    ],
    "安装部署": [
        (
            "Setup_向导",
            [
                "多步向导：选择模式（新装/升级/备份恢复/改配置/Restore FacilityPro）、数据库连接、执行、完成；",
                "写入安装进度并检测数据库健康与 initialized 状态。",
            ],
        ),
        (
            "Setup_数据库",
            [
                "独立数据库初始化与从备份恢复页面（db-init / db-restore）。",
            ],
        ),
        (
            "Setup_采样器校准",
            [
                "采样器校准流程页面，用于现场设备校准记录与参数下发。",
            ],
        ),
    ],
}


def fmt_rules(items: list[str]) -> str:
    lines = ["本功能提供："]
    for i, t in enumerate(items, 1):
        lines.append(f"{i}、{t.rstrip('；;')};")
    return "\n".join(lines)


def write_md():
    DOCS.mkdir(parents=True, exist_ok=True)
    parts = [
        "# 软件需求描述文档",
        "",
        "> 产品：Tauri + Vue 运维管理端（Aventino / Soybean 扩展）",
        "> 说明：本文档按业务大模块划分，子条目采用「模块_功能点」命名。",
        "",
        "---",
        "",
    ]
    for mod, entries in MODULES.items():
        parts.append(f"## {mod}")
        parts.append("")
        for name, rules in entries:
            parts.append(f"### {name}")
            parts.append("")
            parts.append(fmt_rules(rules))
            parts.append("")
    path = DOCS / "软件需求描述.md"
    path.write_text("\n".join(parts), encoding="utf-8")
    return path


def write_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        import subprocess
        import sys

        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font

    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for mod, entries in MODULES.items():
        # Excel sheet name max 31 chars
        sheet_name = mod[:31]
        ws = wb.create_sheet(title=sheet_name)
        ws.append(["功能点", "功能细则描述"])
        for cell in ws[1]:
            cell.font = header_font
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 88
        for name, rules in entries:
            ws.append([name, fmt_rules(rules)])
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
            for cell in row:
                cell.alignment = wrap

    path = DOCS / "软件需求描述.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    md_path = write_md()
    xlsx_path = write_xlsx()
    print(md_path)
    print(xlsx_path)
